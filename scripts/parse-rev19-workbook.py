#!/usr/bin/env python3
"""
Read a filled-in RPS QUOTE TOOL (REV19) workbook and print it as JSON that
scripts/import-rev19-quote.ts can load into Billing.

Only the MATERIAL AND LABOR BREAKDOWN tab is trusted for line items - it is the
sheet the face totals are built from. Header fields come from the same tab, the
project description / ATTN from RP QUOTE TEMPLATE, and the scope row from
COST ESTIMATE when one is filled in.

    python3 scripts/parse-rev19-workbook.py "<workbook>.xlsx" > quote.json
"""
import json, sys
import openpyxl

PINK = {"FFFCE4EC", "FFFFC7CE", "FFF4CCCC"}   # workbook "estimate / needs verified" fills

def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None

def txt(v):
    s = str(v).strip() if v is not None else ""
    return s or None

def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["MATERIAL AND LABOR BREAKDOWN"]

    # ── header + INPUTS block (fixed cells on the REV19 layout) ──────────────
    header = {
        "customer": txt(ws["C5"].value), "site_number": txt(ws["C6"].value), "facility_address": txt(ws["C7"].value),
        "city_state_zip": txt(ws["C8"].value), "proposal_date": txt(ws["C9"].value), "bid_due": txt(ws["C10"].value),
        "project_manager": txt(ws["C11"].value), "construction_manager": txt(ws["C12"].value), "foreman": txt(ws["C13"].value),
        "compiled_by": txt(ws["C14"].value), "csr_number": txt(ws["C15"].value),
        "signer_name": txt(ws["C433"].value), "signer_title": txt(ws["C434"].value),
    }
    inputs = {
        "material_markup_pct": num(ws["D18"].value) or 0, "material_tax_pct": num(ws["D19"].value) or 0,
        "sub_markup_pct": num(ws["D20"].value) or 0, "labor_rate": num(ws["D21"].value) or 0,
        "contingency_pct": num(ws["D22"].value) or 0, "profit_overhead_pct": num(ws["D23"].value) or 0,
        "contingency_flat": 0, "sales_tax_pct": 0,
    }

    # ── find each CATEGORY n block by scanning column B ──────────────────────
    starts = {}
    for r in range(1, ws.max_row + 1):
        b = txt(ws.cell(r, 2).value)
        if b and b.upper().startswith("CATEGORY "):
            try: starts[int(b.split()[1])] = r
            except ValueError: pass
    ends = {c: (starts.get(c + 1) or ws.max_row) for c in starts}

    def flag(r):
        c = ws.cell(r, 3)
        rgb = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type else None
        return "estimate" if rgb in PINK else "ok"

    lines = []
    for cat, r0 in sorted(starts.items()):
        for r in range(r0 + 1, ends[cat]):   # header rows fall out on the numeric checks below
            B, C = txt(ws.cell(r, 2).value), txt(ws.cell(r, 3).value)
            if B and (B.upper().startswith("TOTAL CATEGORY") or B.upper() == "SUBCONTRACTOR MARKUP"): continue
            if C and C.upper().startswith(("TOTAL CATEGORY", "CONSTRUCTION TEAM SUBTOTAL", "SERVICE TECHNICIAN SUBTOTAL")): continue
            if cat == 7:
                men, hrs = num(ws.cell(r, 5).value), num(ws.cell(r, 6).value)
                if not (men and hrs): continue
                lines.append({"section": "basic", "category": 7, "day_label": B, "description": C, "men": men, "hrs_each": hrs,
                              "labor_rate": num(ws.cell(r, 4).value), "crew": "service" if (B or "").upper().startswith("SERVICE") else "construction"})
            elif cat == 8:
                days, techs = num(ws.cell(r, 6).value), num(ws.cell(r, 7).value)
                if not (days and techs): continue
                lines.append({"section": "basic", "category": 8, "day_label": B, "description": C, "unit_cost": num(ws.cell(r, 4).value),
                              "travel_days": days, "techs": techs, "crew": "service" if (B or "").upper().startswith("SERVICE") else "construction",
                              "source_note": txt(ws.cell(r, 11).value)})
            elif cat == 6:
                qty = num(ws.cell(r, 9).value)
                if not qty: continue
                lines.append({"section": "basic", "category": 6, "description": f"{B} - {C}", "unit_cost": num(ws.cell(r, 4).value), "quantity": qty,
                              "source_note": txt(ws.cell(r, 11).value)})
            else:
                qty = num(ws.cell(r, 9).value)
                if not qty or not C: continue
                row = {"section": "basic", "category": cat, "part_number": B, "description": C, "unit_cost": num(ws.cell(r, 4).value), "quantity": qty,
                       "source_note": txt(ws.cell(r, 11).value), "price_flag": flag(r)}
                if cat <= 4:
                    row["freight_per_unit"] = num(ws.cell(r, 7).value) or 0
                else:
                    row["markup_applies"] = (txt(ws.cell(r, 5).value) or "N").upper().startswith("Y")
                lines.append(row)

    face = wb["RP QUOTE TEMPLATE"]
    attn = txt(face["C17"].value)
    description = txt(face["E21"].value)
    scope_rows = []
    if "COST ESTIMATE" in wb.sheetnames:
        ce = wb["COST ESTIMATE"]
        for r in range(19, 21):
            s, d = txt(ce.cell(r, 1).value), txt(ce.cell(r, 3).value)
            if s or d: scope_rows.append({"scope": s or "", "description": d or ""})
    workbook_total = num(ws["J430"].value)

    print(json.dumps({"source_file": path.split("/")[-1], "header": header, "attn": attn, "project_description": description,
                      "scope_rows": scope_rows, "inputs": inputs, "lines": lines, "workbook_total": workbook_total}, indent=2))

if __name__ == "__main__":
    main(sys.argv[1])
